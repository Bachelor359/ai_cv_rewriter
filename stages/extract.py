"""
Stage 1 – Extract CV structure from a DOCX file.

Produces a CVStructure (list of CVSlot) where every addressable paragraph
inside the document gets a sequential slot_id and enough metadata to locate
and overwrite it later in Stage 3.

Handled content locations (in order of slot assignment):
  1. Body paragraphs  (doc.paragraphs — top-level only, no tables/textboxes)
  2. Table cell paragraphs  (all top-level tables; merged cells deduplicated)
  3. Header paragraphs  (per section, skips linked-to-previous)
  4. Footer paragraphs  (per section, skips linked-to-previous)
  5. Text-box paragraphs  (via XML walk of <w:txbxContent>)

Critical design notes:
  - Text is read verbatim; no run consolidation happens here (that's Stage 3).
  - Merged table cells share a <w:tc> element — we deduplicate by element id.
  - Text boxes live inside <w:drawing> or VML shapes; both use <w:txbxContent>.
  - Empty paragraphs are kept as slots so indices stay stable for Stage 3.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from models import CVSlot, CVStructure, SlotLocation


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def extract_cv_structure(docx_path: Path | str) -> CVStructure:
    doc = Document(str(docx_path))
    slots: list[CVSlot] = []
    slot_id = 0

    # --- 1. Body paragraphs ------------------------------------------------
    for para_idx, para in enumerate(doc.paragraphs):
        slots.append(CVSlot(
            slot_id=slot_id,
            text=para.text,
            location=SlotLocation.BODY,
            style_name=_style_name(para),
            indices=[para_idx],
        ))
        slot_id += 1

    # --- 2. Table cell paragraphs ------------------------------------------
    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            seen_cells: set[int] = set()  # deduplicate merged cells by element id
            for col_idx, cell in enumerate(row.cells):
                cell_elem_id = id(cell._tc)
                if cell_elem_id in seen_cells:
                    continue
                seen_cells.add(cell_elem_id)

                for para_idx, para in enumerate(cell.paragraphs):
                    slots.append(CVSlot(
                        slot_id=slot_id,
                        text=para.text,
                        location=SlotLocation.TABLE,
                        style_name=_style_name(para),
                        indices=[table_idx, row_idx, col_idx, para_idx],
                    ))
                    slot_id += 1

    # --- 3 & 4. Header / Footer paragraphs ---------------------------------
    for section_idx, section in enumerate(doc.sections):
        hdr = section.header
        if hdr and not hdr.is_linked_to_previous:
            for para_idx, para in enumerate(hdr.paragraphs):
                slots.append(CVSlot(
                    slot_id=slot_id,
                    text=para.text,
                    location=SlotLocation.HEADER,
                    style_name=_style_name(para),
                    indices=[section_idx, para_idx],
                ))
                slot_id += 1

        ftr = section.footer
        if ftr and not ftr.is_linked_to_previous:
            for para_idx, para in enumerate(ftr.paragraphs):
                slots.append(CVSlot(
                    slot_id=slot_id,
                    text=para.text,
                    location=SlotLocation.FOOTER,
                    style_name=_style_name(para),
                    indices=[section_idx, para_idx],
                ))
                slot_id += 1

    # --- 5. Text-box paragraphs (XML walk) ---------------------------------
    # <w:txbxContent> can be nested inside <w:drawing>/<wp:inline|anchor>
    # or inside VML <v:textbox>. Both are reachable via body.iter().
    for tb_idx, tb_elem in enumerate(doc.element.body.iter(qn("w:txbxContent"))):
        para_elems = tb_elem.findall(f".//{qn('w:p')}")
        for para_idx, p_elem in enumerate(para_elems):
            text = _text_from_p_elem(p_elem)
            slots.append(CVSlot(
                slot_id=slot_id,
                text=text,
                location=SlotLocation.TEXTBOX,
                style_name="TextBox",  # style not trivially available from raw XML
                indices=[tb_idx, para_idx],
            ))
            slot_id += 1

    _mark_continuations(slots)
    return CVStructure(slots=slots)


# ---------------------------------------------------------------------------
# CLI entry-point (python -m stages.extract  or  python stages/extract.py)
# ---------------------------------------------------------------------------

def main() -> None:
    import sys
    from config import CV_PATH, TMP_DIR

    docx_path = Path(sys.argv[1]) if len(sys.argv) > 1 else CV_PATH
    if not docx_path.exists():
        print(f"[extract] CV not found at: {docx_path}")
        sys.exit(1)

    print(f"[extract] Reading: {docx_path}")
    structure = extract_cv_structure(docx_path)

    out_path = TMP_DIR / "cv_structure.json"
    out_path.write_text(structure.model_dump_json(indent=2), encoding="utf-8")

    non_empty = len(structure.non_empty_slots())
    print(f"[extract] {len(structure.slots)} total slots  ({non_empty} non-empty)")
    print(f"[extract] Saved -> {out_path}")

    # Print a preview
    print("\n--- Slot preview (non-empty) ---")
    for s in structure.non_empty_slots()[:20]:
        preview = s.text[:80].replace("\n", " ")
        print(f"  [{s.slot_id:3}] ({s.location.value:8} | {s.style_name[:20]:20}) {preview}")
    if non_empty > 20:
        print(f"  ... and {non_empty - 20} more")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _style_name(para) -> str:
    try:
        return para.style.name if para.style else "Normal"
    except Exception:
        return "Normal"


def _text_from_p_elem(p_elem) -> str:
    """Concatenate all <w:t> text nodes inside a <w:p> element."""
    return "".join(
        t.text or ""
        for t in p_elem.iter(qn("w:t"))
    )


# Words that commonly end a wrapped line but are NOT sentence endings.
# If the previous slot ends with one of these, the next slot is a continuation
# even when it starts with an uppercase letter (e.g. proper noun / tech name).
_WRAP_ENDING_WORDS = frozenset({
    "and", "or", "to", "for", "with", "in", "of", "the", "a", "an",
    "by", "on", "at", "from", "into", "through", "using", "via",
    "as", "is", "are", "was", "were", "be", "been",
})


def _mark_continuations(slots: list) -> None:
    """
    Detect visual line-wrap continuations and set slot.merged_into.

    A slot is a continuation of the preceding slot when ALL of:
      - The preceding slot has non-empty text
      - Both slots share the same location type (don't merge across sections)
      - ANY of these continuation signals:
          a) Current text (stripped) starts with a lowercase letter
          b) Previous text ends with a hyphen  (mid-word line break)
          c) Previous text's last word is a small conjunction/preposition
             (handles "…Spark, and\nAirflow…" where Airflow is uppercase)
    """
    for i in range(1, len(slots)):
        prev = slots[i - 1]
        curr = slots[i]

        prev_stripped = prev.text.strip()
        curr_stripped = curr.text.strip()

        if not prev_stripped or not curr_stripped:
            continue  # either side empty — not a continuation

        if prev.location != curr.location:
            continue  # never merge across location types

        # Determine the primary for the previous slot (follow the chain)
        prev_primary = prev.merged_into if prev.merged_into is not None else prev.slot_id

        # Signal (a): current starts with a lowercase letter
        is_lowercase_start = curr_stripped[0].islower()

        # Signal (b): previous ended with a hyphen (mid-word break)
        is_hyphen_break = prev_stripped.endswith("-") and not prev_stripped.endswith("--")

        # Signal (c): previous last word is a wrap-inducing function word
        last_word = prev_stripped.split()[-1].lower().rstrip(".,;:")
        is_conjunction_end = last_word in _WRAP_ENDING_WORDS

        if is_lowercase_start or is_hyphen_break or is_conjunction_end:
            curr.merged_into = prev_primary


if __name__ == "__main__":
    main()
